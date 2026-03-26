from io import BytesIO
from pathlib import Path
import shutil
import uuid
import zipfile

from openpyxl import load_workbook
import pytest
from pypdf import PdfReader, PdfWriter

import app as app_module
from invoice_helper.ledger import init_ledger_db
from invoice_helper.models import AirlineInvoiceFields, GeneralInvoiceFields, RailwayTicketFields


@pytest.fixture()
def client(monkeypatch):
    app_module.app.config["TESTING"] = True
    temp_root = Path.cwd() / ".pytest_tmp" / str(uuid.uuid4())
    temp_root.mkdir(parents=True, exist_ok=True)
    upload_dir = temp_root / "uploads"
    upload_dir.mkdir(exist_ok=True)
    stats_dedup_dir = temp_root / "stats_dedup"
    stats_dedup_dir.mkdir(exist_ok=True)
    db_path = temp_root / "ledger.sqlite3"
    monkeypatch.setattr(app_module, "UPLOAD_DIR", upload_dir)
    monkeypatch.setattr(app_module, "LEDGER_DB_PATH", db_path)
    monkeypatch.setattr(app_module, "STATS_DEDUP_DIR", stats_dedup_dir)
    init_ledger_db(db_path)
    app_module.RAILWAY_DOCUMENT_STORE.clear()
    app_module.GENERAL_DOCUMENT_STORE.clear()
    app_module.AIRLINE_DOCUMENT_STORE.clear()
    app_module.STATS_DEDUP_TASK_STORE.clear()
    yield app_module.app.test_client()
    shutil.rmtree(temp_root, ignore_errors=True)


def test_upload_preview_and_export_flow(client, monkeypatch):
    def fake_parse(file_name, file_bytes):
        return RailwayTicketFields(
            invoice_number="25429165800000526150",
            issue_date="2025-04-01",
            departure_station="武汉站",
            arrival_station="北京西站",
            departure_datetime="2025-03-31 08:36",
            train_number="G70",
            seat_number="08车05B号",
            amount="623.00",
            passenger_name="李志",
            passenger_id="4307231982****0037",
        )

    monkeypatch.setattr(app_module, "parse_railway_ticket_from_bytes", fake_parse)

    upload_response = client.post(
        "/api/railway/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "ticket.pdf")},
        content_type="multipart/form-data",
    )

    assert upload_response.status_code == 200
    documents = upload_response.get_json()
    assert len(documents) == 1
    assert documents[0]["fields"]["departure_station"] == "武汉站"

    preview_response = client.post(
        "/api/railway/preview-rename",
        json={
            "documentIds": [documents[0]["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{出发站}_{到达站}_{票价}",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )

    assert preview_response.status_code == 200
    preview_payload = preview_response.get_json()
    assert preview_payload[0]["conflictResolvedName"] == "2025-04-01_武汉站_北京西站_623.00"

    export_response = client.post(
        "/api/railway/export",
        json={
            "documentIds": [documents[0]["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{出发站}_{到达站}_{票价}",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )

    assert export_response.status_code == 200
    assert export_response.mimetype == "application/zip"

    excel_response = client.post(
        "/api/railway/export-excel",
        json={
            "documentIds": [documents[0]["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{出发站}_{到达站}_{票价}",
                "dateFormat": "YYYY-MM-DD",
            },
            "excelColumns": ["invoice_number", "departure_station", "amount"],
        },
    )

    assert excel_response.status_code == 200
    assert excel_response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(excel_response.data))
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    row = [cell.value for cell in worksheet[2]]
    assert headers == ["原文件名", "发票号码", "出发站", "票价"]
    assert row[0] == "ticket.pdf"
    assert row[1:] == ["25429165800000526150", "武汉站", "623.00"]


def test_delete_and_clear_documents(client, monkeypatch):
    def fake_parse(file_name, file_bytes):
        return RailwayTicketFields(
            invoice_number="25429165800000526150",
            issue_date="2025-04-01",
            departure_station="武汉站",
            arrival_station="北京西站",
            departure_datetime="2025-03-31 08:36",
            amount="623.00",
        )

    monkeypatch.setattr(app_module, "parse_railway_ticket_from_bytes", fake_parse)

    first = client.post(
        "/api/railway/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 first"), "first.pdf")},
        content_type="multipart/form-data",
    ).get_json()[0]
    second = client.post(
        "/api/railway/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 second"), "second.pdf")},
        content_type="multipart/form-data",
    ).get_json()[0]

    delete_response = client.delete(f"/api/railway/documents/{first['id']}")
    assert delete_response.status_code == 200
    assert first["id"] not in app_module.RAILWAY_DOCUMENT_STORE
    assert second["id"] in app_module.RAILWAY_DOCUMENT_STORE

    clear_response = client.delete("/api/railway/documents")
    assert clear_response.status_code == 200
    assert app_module.RAILWAY_DOCUMENT_STORE == {}


def test_general_invoice_flow(client, monkeypatch):
    def fake_parse(file_name, file_bytes):
        return GeneralInvoiceFields(
            invoice_type="电子发票（普通发票）",
            invoice_number="25312000000002446025",
            issue_date="2025-01-03",
            buyer_name="广东工业大学",
            seller_name="上海淳大酒店投资管理有限公司",
            total_amount="813.64",
        )

    monkeypatch.setattr(app_module, "parse_general_invoice_from_bytes", fake_parse)

    upload_response = client.post(
        "/api/general-invoice/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "hotel.pdf")},
        content_type="multipart/form-data",
    )

    assert upload_response.status_code == 200
    document = upload_response.get_json()[0]
    assert document["fields"]["seller_name"] == "上海淳大酒店投资管理有限公司"

    preview_response = client.post(
        "/api/general-invoice/preview-rename",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{销售方名称}_{价税合计}",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )
    assert preview_response.status_code == 200
    assert preview_response.get_json()[0]["conflictResolvedName"] == "2025-01-03_上海淳大酒店投资管理有限公司_813.64"

    excel_response = client.post(
        "/api/general-invoice/export-excel",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{销售方名称}_{价税合计}",
                "dateFormat": "YYYY-MM-DD",
            },
            "excelColumns": [],
        },
    )

    assert excel_response.status_code == 200
    workbook = load_workbook(BytesIO(excel_response.data))
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    row = [cell.value for cell in worksheet[2]]
    assert headers[0] == "原文件名"
    assert "购买方名称" in headers
    assert "销售方名称" in headers
    assert "解析状态" not in headers
    assert row[headers.index("销售方名称")] == "上海淳大酒店投资管理有限公司"


def test_airline_invoice_flow(client, monkeypatch):
    def fake_parse(file_name, file_bytes):
        return AirlineInvoiceFields(
            invoice_number="26112000001054174981",
            issue_date="2026-03-18",
            departure_airport="北京大兴",
            arrival_airport="广州",
            flight_number="JD5921",
            cabin_class="经济舱 Q舱",
            departure_time="2025-09-15",
            amount="807.34",
            total_amount="930.00",
            passenger_name="李志",
            passenger_id="430723******120037",
        )

    monkeypatch.setattr(app_module, "parse_airline_invoice_from_bytes", fake_parse)

    upload_response = client.post(
        "/api/airline/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "airline.pdf")},
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200
    document = upload_response.get_json()[0]
    assert document["fields"]["flight_number"] == "JD5921"

    preview_response = client.post(
        "/api/airline/preview-rename",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{起飞机场}_{着陆机场}_{航班号}_{价税合计}",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )
    assert preview_response.status_code == 200
    assert preview_response.get_json()[0]["conflictResolvedName"] == "2026-03-18_北京大兴_广州_JD5921_930.00"

    excel_response = client.post(
        "/api/airline/export-excel",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "template",
                "template": "{开票日期}_{起飞机场}_{着陆机场}_{航班号}_{价税合计}",
                "dateFormat": "YYYY-MM-DD",
            },
            "excelColumns": ["flight_number", "passenger_name"],
        },
    )

    assert excel_response.status_code == 200
    workbook = load_workbook(BytesIO(excel_response.data))
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    row = [cell.value for cell in worksheet[2]]
    assert headers == ["原文件名", "航班号", "乘机人姓名"]
    assert row == ["airline.pdf", "JD5921", "李志"]


def test_export_excel_requires_successful_documents(client):
    upload_response = client.post(
        "/api/railway/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "broken.pdf")},
        content_type="multipart/form-data",
    )

    assert upload_response.status_code == 200
    document = upload_response.get_json()[0]

    excel_response = client.post(
        "/api/railway/export-excel",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {"mode": "tokens"},
            "excelColumns": ["invoice_number"],
        },
    )

    assert excel_response.status_code == 400
    assert excel_response.get_json()["message"] == "没有可导出的成功解析票据。"


def test_split_folder_preview_and_export_flow(client, monkeypatch):
    def fake_parse(file_name, file_bytes):
        return RailwayTicketFields(
            invoice_number="25429165800000526150",
            issue_date="2025-04-01",
            departure_station="武汉站",
            arrival_station="北京西站",
            departure_datetime="2025-03-31 08:36",
            train_number="G70",
            seat_number="08车05B号",
            amount="623.00",
            passenger_name="李志",
        )

    monkeypatch.setattr(app_module, "parse_railway_ticket_from_bytes", fake_parse)

    upload_response = client.post(
        "/api/railway/upload-and-parse",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "ticket.pdf")},
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200
    document = upload_response.get_json()[0]

    preview_response = client.post(
        "/api/railway/preview-split-folder",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "tokens",
                "tokens": [
                    {"type": "field", "value": "issue_date"},
                    {"type": "field", "value": "departure_station"},
                    {"type": "field", "value": "arrival_station"},
                    {"type": "field", "value": "amount"},
                ],
                "separator": "/",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )
    assert preview_response.status_code == 200
    preview_payload = preview_response.get_json()[0]
    assert preview_payload["folderPath"] == "2025-04-01/武汉站/北京西站/623.00"
    assert preview_payload["newFileName"] == "2025-04-01_武汉站_北京西站_623.00.pdf"
    assert preview_payload["fullOutputPath"] == "2025-04-01/武汉站/北京西站/623.00/2025-04-01_武汉站_北京西站_623.00.pdf"

    export_response = client.post(
        "/api/railway/export-split-folder",
        json={
            "documentIds": [document["id"]],
            "ruleConfig": {
                "mode": "tokens",
                "tokens": [
                    {"type": "field", "value": "issue_date"},
                    {"type": "field", "value": "departure_station"},
                    {"type": "field", "value": "arrival_station"},
                    {"type": "field", "value": "amount"},
                ],
                "separator": "/",
                "dateFormat": "YYYY-MM-DD",
            },
        },
    )
    assert export_response.status_code == 200
    assert export_response.mimetype == "application/zip"

    with zipfile.ZipFile(BytesIO(export_response.data)) as archive:
        names = archive.namelist()
    assert names == ["2025-04-01/武汉站/北京西站/623.00/2025-04-01_武汉站_北京西站_623.00.pdf"]


def test_ledger_upload_list_detail_download_and_delete(client, monkeypatch):
    def fake_detect(file_name, file_bytes):
        if "rail" in file_name:
            return (
                "railway",
                RailwayTicketFields(
                    invoice_number="25429165848000965552",
                    issue_date="2025-03-31",
                    departure_station="广州南站",
                    arrival_station="长沙南站",
                    departure_datetime="2025-03-30 16:21",
                    train_number="G810",
                    amount="314.00",
                    passenger_name="李志",
                ).to_dict(),
                "铁路电子客票",
            )
        if "air" in file_name:
            return (
                "airline",
                AirlineInvoiceFields(
                    invoice_number="26112000001054174981",
                    issue_date="2025-09-15",
                    departure_airport="北京大兴",
                    arrival_airport="广州",
                    flight_number="JD5921",
                    cabin_class="经济舱 Q舱",
                    departure_time="2025-09-15",
                    total_amount="930.00",
                    passenger_name="李志",
                ).to_dict(),
                "北京首都航空有限公司",
            )
        if "hotel" in file_name:
            return (
                "general",
                GeneralInvoiceFields(
                    invoice_type="电子发票（普通发票）",
                    invoice_number="25312000000002446025",
                    issue_date="2025-01-08",
                    buyer_name="广东工业大学",
                    seller_name="北京大小酒店有限公司雅乐轩饭店",
                    total_amount="629.64",
                    remarks="*住宿服务*住宿服务",
                ).to_dict(),
                "*住宿服务*住宿服务",
            )
        raise ValueError("未识别为支持的发票类型。")

    monkeypatch.setattr(app_module, "detect_invoice_for_ledger", fake_detect)

    upload_response = client.post(
        "/api/ledger/upload-and-parse",
        data={
            "files": [
                (BytesIO(b"%PDF rail"), "rail-ticket.pdf"),
                (BytesIO(b"%PDF air"), "air-ticket.pdf"),
                (BytesIO(b"%PDF hotel"), "hotel-ticket.pdf"),
                (BytesIO(b"%PDF bad"), "bad-ticket.pdf"),
            ]
        },
        content_type="multipart/form-data",
    )

    assert upload_response.status_code == 200
    payload = upload_response.get_json()
    assert len(payload["created"]) == 3
    assert len(payload["failed"]) == 1

    list_response = client.get("/api/ledger/list")
    assert list_response.status_code == 200
    entries = list_response.get_json()
    assert len(entries) == 3
    assert {entry["expenseType"] for entry in entries} == {"交通", "住宿"}

    transport_response = client.get("/api/ledger/list?expense_types=交通")
    assert transport_response.status_code == 200
    transport_entries = transport_response.get_json()
    assert len(transport_entries) == 2
    assert all(entry["expenseType"] == "交通" for entry in transport_entries)

    hotel_entry = next(entry for entry in entries if entry["expenseType"] == "住宿")
    detail_response = client.get(f"/api/ledger/{hotel_entry['id']}")
    assert detail_response.status_code == 200
    detail = detail_response.get_json()
    assert detail["title"] == "北京大小酒店有限公司雅乐轩饭店"
    assert detail["infoItems"][3]["label"] == "项目名称"

    preview_response = client.get(f"/api/ledger/{hotel_entry['id']}/preview")
    assert preview_response.status_code == 200

    single_download = client.get(f"/api/ledger/download?ids={hotel_entry['id']}")
    assert single_download.status_code == 200

    multi_download = client.get("/api/ledger/download?ids=" + ",".join(entry["id"] for entry in entries))
    assert multi_download.status_code == 200
    assert multi_download.mimetype == "application/zip"

    delete_response = client.delete("/api/ledger", json={"ids": [hotel_entry["id"]]})
    assert delete_response.status_code == 200
    after_delete = client.get("/api/ledger/list").get_json()
    assert len(after_delete) == 2


def test_ledger_stats_summary_charts_and_export(client, monkeypatch):
    def fake_detect(file_name, file_bytes):
        if "rail" in file_name:
            return (
                "railway",
                RailwayTicketFields(
                    invoice_number="25429165848000965552",
                    issue_date="2025-03-31",
                    departure_station="广州南站",
                    arrival_station="长沙南站",
                    departure_datetime="2025-03-30 16:21",
                    train_number="G810",
                    amount="314.00",
                    passenger_name="李志",
                ).to_dict(),
                "铁路电子客票",
            )
        if "rail-dup" in file_name:
            return (
                "railway",
                RailwayTicketFields(
                    invoice_number="25429165848000965552",
                    issue_date="2025-03-31",
                    departure_station="广州南站",
                    arrival_station="长沙南站",
                    departure_datetime="2025-03-30 16:21",
                    train_number="G810",
                    amount="314.00",
                    passenger_name="李志",
                ).to_dict(),
                "铁路电子客票",
            )
        if "air" in file_name:
            return (
                "airline",
                AirlineInvoiceFields(
                    invoice_number="26112000001054174981",
                    issue_date="2025-09-15",
                    departure_airport="北京大兴",
                    arrival_airport="广州",
                    flight_number="JD5921",
                    cabin_class="经济舱 Q舱",
                    departure_time="2025-09-15",
                    total_amount="930.00",
                    passenger_name="李志",
                ).to_dict(),
                "北京首都航空有限公司",
            )
        return (
            "general",
            GeneralInvoiceFields(
                invoice_type="电子发票（普通发票）",
                invoice_number="25312000000002446025",
                issue_date="2025-01-08",
                buyer_name="广东工业大学",
                seller_name="北京大小酒店有限公司雅乐轩饭店",
                total_amount="629.64",
                remarks="*住宿服务*住宿服务",
            ).to_dict(),
            "*住宿服务*住宿服务",
        )

    monkeypatch.setattr(app_module, "detect_invoice_for_ledger", fake_detect)

    upload_response = client.post(
        "/api/ledger/upload-and-parse",
        data={
            "files": [
                (BytesIO(b"%PDF rail"), "rail-ticket.pdf"),
                (BytesIO(b"%PDF rail dup"), "rail-dup-ticket.pdf"),
                (BytesIO(b"%PDF air"), "air-ticket.pdf"),
                (BytesIO(b"%PDF hotel"), "hotel-ticket.pdf"),
            ]
        },
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200

    summary_response = client.get("/api/ledger/stats/summary")
    assert summary_response.status_code == 200
    summary = summary_response.get_json()
    assert summary["invoiceCount"] == 4
    assert summary["uniqueInvoiceCount"] == 3
    assert summary["duplicateCount"] == 1
    assert summary["totalAmount"] == "2187.64"

    charts_response = client.get("/api/ledger/stats/charts")
    assert charts_response.status_code == 200
    charts_payload = charts_response.get_json()
    assert "amountTrend" in charts_payload["charts"]
    assert "invoiceCategoryAmount" in charts_payload["charts"]
    assert charts_payload["charts"]["invoiceCategoryAmount"]["labels"] == ["常规数电发票", "铁路电子客票", "航空电子客票"]
    assert charts_payload["charts"]["invoiceCategoryAmount"]["series"] == [629.64, 628.0, 930.0]

    table_response = client.get("/api/ledger/stats/table")
    assert table_response.status_code == 200
    table_payload = table_response.get_json()
    assert len(table_payload["table"]) == 1
    assert table_payload["table"][0]["duplicateStatus"] == "重复发票"

    export_response = client.get("/api/ledger/stats/export")
    assert export_response.status_code == 200
    workbook = load_workbook(BytesIO(export_response.data))
    assert workbook.sheetnames == ["汇总概览", "明细数据"]
    assert workbook["汇总概览"]["A2"].value == "发票总数"
    assert workbook["明细数据"]["H2"].value == "重复发票"


def test_merge_print_upload_build_download_and_export_list(client):
    pdf_buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(pdf_buffer)
    pdf_bytes = pdf_buffer.getvalue()

    upload_response = client.post(
        "/api/merge-print/upload",
        data={"files": [(BytesIO(pdf_bytes), "first.pdf"), (BytesIO(pdf_bytes), "second.pdf")]},
        content_type="multipart/form-data",
    )

    assert upload_response.status_code == 200
    upload_payload = upload_response.get_json()
    assert upload_payload["inputType"] == "pdf"
    assert len(upload_payload["items"]) == 2

    build_response = client.post(
        "/api/merge-print/build",
        json={
            "taskId": upload_payload["taskId"],
            "itemIds": [item["id"] for item in upload_payload["items"]],
            "layoutMode": "double",
            "showDivider": True,
            "listPlacement": "omit",
        },
    )
    assert build_response.status_code == 200
    build_payload = build_response.get_json()
    assert build_payload["pageCount"] == 1
    assert build_payload["stats"]["invoiceCount"] == 2

    preview_response = client.get(build_payload["previewUrl"])
    assert preview_response.status_code == 200
    assert preview_response.mimetype == "application/pdf"
    preview_reader = PdfReader(BytesIO(preview_response.data))
    assert len(preview_reader.pages) == 1
    assert round(float(preview_reader.pages[0].mediabox.width), 2) == 300.0
    assert round(float(preview_reader.pages[0].mediabox.height), 2) == 816.0

    download_response = client.get(build_payload["downloadUrl"])
    assert download_response.status_code == 200
    assert download_response.mimetype == "application/pdf"

    export_list_response = client.post(
        "/api/merge-print/export-list",
        json={
            "taskId": upload_payload["taskId"],
            "itemIds": [item["id"] for item in upload_payload["items"]],
        },
    )
    assert export_list_response.status_code == 200
    workbook = load_workbook(BytesIO(export_list_response.data))
    worksheet = workbook.active
    assert worksheet["A2"].value == 1
    assert worksheet["B2"].value == "first.pdf"


def test_merge_print_blocks_mixed_uploads(client):
    pdf_buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(pdf_buffer)
    pdf_bytes = pdf_buffer.getvalue()

    ofd_buffer = BytesIO()
    with zipfile.ZipFile(ofd_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("Doc_0/Pages/Page_0/Content.xml", "<ofd><TextObject Value='测试OFD内容'/></ofd>")

    first_response = client.post(
        "/api/merge-print/upload",
        data={"files": (BytesIO(pdf_bytes), "first.pdf")},
        content_type="multipart/form-data",
    )
    assert first_response.status_code == 200
    task_id = first_response.get_json()["taskId"]

    mixed_response = client.post(
        "/api/merge-print/upload",
        data={"taskId": task_id, "files": (BytesIO(ofd_buffer.getvalue()), "second.ofd")},
        content_type="multipart/form-data",
    )
    assert mixed_response.status_code == 400
    assert mixed_response.get_json()["message"] == "同一次合并任务不能同时上传 PDF 和 OFD 文件。"


def test_stats_dedup_analyze_and_export(client, monkeypatch):
    def fake_railway_parse(file_name, file_bytes):
        if "ticket" not in file_name:
            raise ValueError("not railway")
        return RailwayTicketFields(
            invoice_number="25429165800000526150",
            issue_date="2025-04-01",
            departure_station="武汉站",
            arrival_station="北京西站",
            departure_datetime="2025-03-31 08:36",
            amount="623.00",
            passenger_name="李志",
        )

    def fake_airline_parse(file_name, file_bytes):
        if "airline" not in file_name:
            raise ValueError("not airline")
        return AirlineInvoiceFields(
            invoice_number="26112000001054174981",
            issue_date="2026-03-18",
            departure_airport="北京大兴",
            arrival_airport="广州",
            flight_number="JD5921",
            amount="807.34",
            total_amount="930.00",
            passenger_name="李志",
        )

    def fake_general_parse(file_name, file_bytes):
        if "hotel" not in file_name:
            raise ValueError("not general")
        return GeneralInvoiceFields(
            invoice_type="电子发票（普通发票）",
            invoice_number="25312000000002446025",
            issue_date="2025-01-03",
            buyer_name="广东工业大学",
            seller_name="上海淳大酒店投资管理有限公司",
            amount="767.58",
            tax_amount="46.06",
            total_amount="813.64",
        )

    monkeypatch.setattr(app_module, "parse_railway_ticket_from_bytes", fake_railway_parse)
    monkeypatch.setattr(app_module, "parse_airline_invoice_from_bytes", fake_airline_parse)
    monkeypatch.setattr(app_module, "parse_general_invoice_from_bytes", fake_general_parse)

    upload_response = client.post(
        "/api/stats-dedup/upload",
        data={
            "files": [
                (BytesIO(b"%PDF rail"), "ticket.pdf"),
                (BytesIO(b"%PDF rail2"), "ticket-duplicate.pdf"),
                (BytesIO(b"%PDF airline"), "airline.pdf"),
                (BytesIO(b"%PDF hotel"), "hotel.pdf"),
            ]
        },
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200
    task = upload_response.get_json()
    assert len(task["items"]) == 4

    analyze_response = client.post("/api/stats-dedup/analyze", json={"taskId": task["taskId"]})
    assert analyze_response.status_code == 200
    payload = analyze_response.get_json()
    assert payload["summary"]["duplicateCount"] == 1
    assert payload["summary"]["successCount"] == 3
    assert payload["summary"]["totalAmount"] == "2366.64"
    statuses = {item["originalName"]: item["dedupStatus"] for item in payload["items"]}
    assert statuses["ticket.pdf"] == "统计完成"
    assert statuses["ticket-duplicate.pdf"] == "重复发票"
    assert statuses["airline.pdf"] == "统计完成"
    assert statuses["hotel.pdf"] == "统计完成"

    export_response = client.post("/api/stats-dedup/export", json={"taskId": task["taskId"]})
    assert export_response.status_code == 200
    export_payload = export_response.get_json()
    download_response = client.get(export_payload["downloadUrl"])
    assert download_response.status_code == 200
    workbook = load_workbook(BytesIO(download_response.data))
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    assert headers == [
        "序号",
        "文件名称",
        "票种",
        "发票号码",
        "开票日期",
        "金额",
        "税额",
        "价税合计",
        "是否重复",
        "重复组",
        "统计状态",
        "错误信息",
    ]
    assert worksheet.max_row == 5
    assert worksheet["I3"].value == "是"


def test_stats_dedup_rejects_non_pdf(client):
    upload_response = client.post(
        "/api/stats-dedup/upload",
        data={"files": (BytesIO(b"fake ofd"), "ticket.ofd")},
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 400
    assert upload_response.get_json()["message"] == "当前模块仅支持 PDF 发票。"


def test_stats_dedup_parse_failure_is_preserved(client, monkeypatch):
    monkeypatch.setattr(app_module, "parse_railway_ticket_from_bytes", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("x")))
    monkeypatch.setattr(app_module, "parse_airline_invoice_from_bytes", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("y")))
    monkeypatch.setattr(app_module, "parse_general_invoice_from_bytes", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("z")))

    upload_response = client.post(
        "/api/stats-dedup/upload",
        data={"files": (BytesIO(b"%PDF-1.4 fake"), "broken.pdf")},
        content_type="multipart/form-data",
    )
    assert upload_response.status_code == 200
    task = upload_response.get_json()

    analyze_response = client.post("/api/stats-dedup/analyze", json={"taskId": task["taskId"]})
    assert analyze_response.status_code == 200
    item = analyze_response.get_json()["items"][0]
    assert item["dedupStatus"] == "解析失败"
    assert item["error"] == "未识别为支持统计的 PDF 发票。"
